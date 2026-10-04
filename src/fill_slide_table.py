"""Fill the "Metric | Paper | Ours" table on the results slide from an experiment's summary JSON.

Usage:
    python src/fill_slide_table.py                       # A0 → slide 18 of the filled deck
    python src/fill_slide_table.py --summary results/A0_baseline/A0_summary.json --slide 18

Only run.text is edited, so the deck's fonts and colours stay as they are.
"""
import argparse
import json

from pptx import Presentation

ROWS = {  # slide row label -> (section in the summary JSON, metric key)
    "Mean fold accuracy":   ("mean_fold", "accuracy"),
    "Ensemble accuracy":    ("ensemble", "accuracy"),
    "Recall (sensitivity)": ("ensemble", "recall"),
    "Precision":            ("ensemble", "precision"),
    "F1-score":             ("ensemble", "f1"),
    "ROC-AUC":              ("ensemble", "auc"),
    "MCC":                  ("ensemble", "mcc"),
}


def fmt(label, value):
    # MCC is stored ×100 in the JSON; the slide shows it as a coefficient
    return f"{value / 100:.2f}" if label == "MCC" else f"{value:.2f} %"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--summary", default="results/A0_baseline/A0_summary.json")
    ap.add_argument("--deck", default="slides/Malaria_Phase2_Baseline_Hypothesis_filled.pptx")
    ap.add_argument("--slide", type=int, default=18)
    args = ap.parse_args()

    summary = json.load(open(args.summary))
    deck = Presentation(args.deck)
    slide = deck.slides[args.slide - 1]
    table = next(sh.table for sh in slide.shapes
                 if sh.has_table and [c.text for c in sh.table.rows[0].cells] == ["Metric", "Paper", "Ours"])

    for row in list(table.rows)[1:]:
        label = row.cells[0].text.strip()
        if label not in ROWS:
            continue
        section, key = ROWS[label]
        runs = [r for p in row.cells[2].text_frame.paragraphs for r in p.runs]
        runs[0].text = fmt(label, summary[section][key])
        for r in runs[1:]:
            r.text = ""
        print(f"{label:22s} {runs[0].text}")

    deck.save(args.deck)
    print(f"Saved {args.deck}")


if __name__ == "__main__":
    main()
